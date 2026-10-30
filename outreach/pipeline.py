"""Pipeline state: merge worker output, apply verifier verdicts, report progress, export.

  python -m outreach.pipeline merge          # data/people/*.jsonl -> data/people.jsonl (validated, deduped)
  python -m outreach.pipeline verdicts       # data/verdicts/*.jsonl -> verified flags + trust audits
  python -m outreach.pipeline status
  python -m outreach.pipeline export         # data/export/professors.xlsx + .csv

A record is usable for outreach only if: schema valid, validate.py clean, verifier verdict pass.
"""
from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

from .config import ROOT
from .schema import Person
from .trust import record_audit
from .validate import check

DATA = ROOT / "data"


def _rows(paths) -> list[dict]:
    out = []
    for p in paths:
        for l in Path(p).read_text().splitlines():
            if l.strip() and not l.startswith('{"_'):
                out.append(json.loads(l))
    return out


def merge() -> None:
    rows = _rows(sorted((DATA / "people").glob("*.jsonl")))
    by_id: dict[str, dict] = {}
    for r in rows:  # later files win (re-runs supersede)
        by_id[r["id"]] = r
    probs = check(list(by_id.values()), DATA / "snapshots")
    bad = Counter(pid for pid, _ in probs)
    status = json.loads((DATA / "status.json").read_text()) if (DATA / "status.json").exists() else {}
    with (DATA / "people.jsonl").open("w") as f:
        for pid, r in by_id.items():
            f.write(json.dumps(r) + "\n")
            st = status.setdefault(pid, {})
            st["validate_ok"] = pid not in bad
    (DATA / "status.json").write_text(json.dumps(status, indent=1))
    (DATA / "validate_report.tsv").write_text("".join(f"{a}\t{b}\n" for a, b in probs))
    print(f"{len(by_id)} people, {len(bad)} with validator problems (data/validate_report.tsv)")


def verdicts() -> None:
    status = json.loads((DATA / "status.json").read_text())
    for v in _rows(sorted((DATA / "verdicts").glob("*.jsonl"))):
        st = status.setdefault(v["id"], {})
        if st.get("verdict_seen") == json.dumps(v, sort_keys=True):
            continue
        st["verified"] = v["verdict"] == "pass"
        st["verifier_errors"] = v.get("field_errors", [])
        st["verdict_seen"] = json.dumps(v, sort_keys=True)
        record_audit("profile", v["id"], v["verdict"] == "pass", json.dumps(v.get("field_errors", []))[:300])
    (DATA / "status.json").write_text(json.dumps(status, indent=1))
    print("verdicts applied")


def status() -> None:
    people = [Person.model_validate(r) for r in _rows([DATA / "people.jsonl"])]
    st = json.loads((DATA / "status.json").read_text()) if (DATA / "status.json").exists() else {}
    c = Counter()
    for p in people:
        s = st.get(p.id, {})
        c["total"] += 1
        c[f"tier {p.tier}"] += 1
        c["early career"] += p.early_career
        c["has email"] += p.email is not None
        c["validated"] += bool(s.get("validate_ok"))
        c["verified"] += bool(s.get("verified"))
        c["outreach-ready (A, email, validated, verified)"] += (p.tier == "A" and p.email is not None
                                                                 and s.get("validate_ok") and s.get("verified"))
    by_inst = Counter(p.institution for p in people)
    for k, v in c.items():
        print(f"{k:48s} {v}")
    print("\nper institution:", dict(by_inst.most_common()))


def export() -> None:
    people = [Person.model_validate(r) for r in _rows([DATA / "people.jsonl"])]
    st = json.loads((DATA / "status.json").read_text()) if (DATA / "status.json").exists() else {}
    cols = ["Name", "University/Institute", "Department", "Position", "Early career", "Start Year", "Email",
            "Email source", "Personal Website", "Lab Website", "Google Scholar", "Research Keywords",
            "Recent Papers", "PhD Recruiting", "Recruiting evidence", "AI relevance", "Tier",
            "Independent PI", "University appointment / PhD path", "Verified", "Notes"]
    out = DATA / "export"
    out.mkdir(exist_ok=True)
    rows = []
    for p in sorted(people, key=lambda p: (p.tier, not p.early_career, p.institution, p.name)):
        s = st.get(p.id, {})
        rows.append([
            p.name, p.institution, "; ".join(p.departments), p.position.value, "yes" if p.early_career else "",
            p.start_year.value if p.start_year else "", p.email.value if p.email else "Not listed",
            str(p.email.evidence.source_url) if p.email else "", str(p.homepage or ""), str(p.lab_website or ""),
            str(p.google_scholar or ""), "; ".join(p.keywords),
            "\n".join(f"{x.title} ({x.year}) {x.url}" for x in p.recent_papers), p.recruiting.status.value,
            p.recruiting.evidence.quote if p.recruiting.evidence else "", p.relevance.value, p.tier,
            ("yes" if p.institute.independent_pi else "no") if p.institute else "",
            ((p.institute.university_appointment.value if p.institute.university_appointment else "") + " " + (p.institute.phd_path or "")).strip() if p.institute else "",
            "yes" if s.get("verified") else ("FAILED" if s.get("verified") is False else "pending"), p.notes,
        ])
    with (out / "professors.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Professors"
        ws.append(cols)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        widths = [22, 26, 26, 30, 8, 8, 28, 30, 30, 30, 30, 40, 60, 12, 40, 10, 6, 8, 30, 10, 50]
        for i, wdt in enumerate(widths):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = wdt
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(out / "professors.xlsx")
    except ImportError:
        pass
    print(f"exported {len(rows)} rows to {out}")


if __name__ == "__main__":
    {"merge": merge, "verdicts": verdicts, "status": status, "export": export}[sys.argv[1]]()
