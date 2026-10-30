"""Who to email first: young, relevant, with an email; Chinese / Asian surnames first, then fit.

  python -m outreach.shortlist [--no-cs]   # -> data/export/shortlist[_noncs].csv + .xlsx, prints a summary
"""
from __future__ import annotations

import csv
import json
from collections import Counter

from .config import ROOT
from .surnames import PRIORITY

DATA = ROOT / "data"
RANK = {"incoming_assistant": 0, "assistant": 0, "associate": 1, "research_faculty": 2, "pi_from_papers": 2}


def bad_emails() -> dict[str, str]:
    """ids whose email failed validate.py (non-WARN): the address is not trusted, the person needs a lookup."""
    f = DATA / "validate_report.tsv"
    out = {}
    for l in (f.read_text().splitlines() if f.exists() else []):
        pid, _, msg = l.partition("\t")
        if msg.startswith("email") and not msg.startswith("WARN"):
            out[pid] = msg
    return out


def rows(skip_cs: bool = False) -> list[dict]:
    seen, out = set(), []
    bad = bad_emails()
    allp = [{**json.loads(l), "_cs": f.name.startswith(("csr-", "cmu-lti")), "_stem": f.stem}
            for f in sorted((DATA / "people").glob("*.jsonl")) for l in f.read_text().splitlines() if l.strip()]
    allp.sort(key=lambda p: (p["email"] is None, -(p.get("fit") or 0)))   # duplicates: keep the copy with an email
    for p in allp:
        keys = {(p["name"].lower(), p["institution"])} | {k for k in (
            (p["email"] or {}).get("value"), (p.get("homepage") or "").rstrip("/").lower(), p.get("google_scholar")) if k}
        if skip_cs and p["_cs"]:
            continue
        if keys & seen or p["relevance"] == "no" or p.get("moved") or p["position_class"] not in RANK:
            continue
        seen |= keys
        if p["id"] in bad and p["email"]:
            p = {**p, "email": None, "notes": f"{p.get('notes', '')}; email rejected: {bad[p['id']]}"}
        out.append(p)
    out.sort(key=lambda p: (PRIORITY[p.get("heritage")], -(p.get("fit") or 0), RANK[p["position_class"]],
                            p["email"] is None, p["name"]))
    return out


def targets() -> dict:
    import yaml
    f = ROOT / "config" / "targets.yaml"
    return yaml.safe_load(f.read_text()) if f.exists() else {}


def main() -> None:
    import sys
    t = targets() if "--targets" in sys.argv else {}
    skip_cs = "--no-cs" in sys.argv or t.get("include_cs") is False
    ps = rows(skip_cs)
    if t:
        ex = set(t.get("exclude_schools", []))
        ps = [p for p in ps if p["institution"] not in ex and (not t.get("heritage") or p.get("heritage") == t["heritage"])]
        focus = set(t.get("focus_schools", []))
        ps.sort(key=lambda p: (p["institution"] not in focus, p["email"] is None, -(p.get("fit") or 0)))
    cols = ["Heritage", "Fit", "Matches project", "Name", "Institution", "Department", "Position", "Email",
            "Recruiting", "Keywords", "Paper 1", "Homepage", "Tier", "Check"]
    table = [[p.get("heritage") or "", p.get("fit") or "", ",".join(p.get("fit_projects") or []), p["name"], p["institution"],
              "; ".join(p["departments"]), p["position"]["value"], (p["email"] or {}).get("value", ""),
              p["recruiting"]["status"], "; ".join(p["keywords"][:5]),
              (f"{p['recent_papers'][0]['title']} ({p['recent_papers'][0]['year']})" if p["recent_papers"] else ""),
              p.get("homepage") or "", p["tier"], "human check" if "human check" in p.get("notes", "") else ""] for p in ps]
    out = DATA / "export"
    out.mkdir(parents=True, exist_ok=True)
    name = "batch1" if t else ("shortlist_noncs" if skip_cs else "shortlist")
    with (out / f"{name}.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(table)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Shortlist"
        ws.append(cols)
        for r in table:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        for col, wdt in zip("ABCDEFGHIJKLMN", [10, 5, 14, 22, 26, 24, 26, 28, 10, 50, 60, 30, 5, 10]):
            ws.column_dimensions[col].width = wdt
        ws2 = wb.create_sheet("Need email (look up)")
        ws2.append(["Fit", "Name", "Institution", "Department", "Position", "Homepage", "Why"])
        for p in ps:
            if not p["email"] and (p.get("fit") or 0) >= 3:
                ws2.append([p.get("fit"), p["name"], p["institution"], "; ".join(p["departments"]), p["position"]["value"],
                            p.get("homepage") or "", (p.get("notes") or "")[-120:]])
        wb.save(out / f"{name}.xlsx")
    except ImportError:
        pass
    with_email = [p for p in ps if p["email"]]
    print(f"{len(ps)} candidates, {len(with_email)} with email; fit>=4 with email: {sum((p.get('fit') or 0) >= 4 for p in with_email)}")
    print("heritage:", dict(Counter(p.get("heritage") for p in with_email)))


if __name__ == "__main__":
    main()
